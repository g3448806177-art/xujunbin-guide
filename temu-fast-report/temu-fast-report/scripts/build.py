#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""TEMU 快速报单 - build 阶段：根据 matches.json 生成报单Excel。

用法:
  python3 build.py --records <records.json> --matches <matches.json> [--output <输出.xlsx>]

matches.json 格式（reporter 可省略，默认使用 prepare 时确定的报单名）:
{
  "date": "0909",
  "matched": {
    "<sourceRow>": {
      "webNo": "Q2", "webColor": "黑色", "cost": 25,
      "shop": "创想鞋业",
      "url": "https://chuang168.sooxie.com/detail/2455431"
    }
  },
  "unmatched": [16]
}

报单格式(严格对齐 assets/采购单模版.xlsx，写死):
  第1行分区标题: A1:L1合并"采购" / N1:S1合并"网购"，宋体系列20号，M列留空分隔
  第2行表头(黄色，左侧除K2商品属性外18号，K2和右侧20号):
    左侧 A-L: 码数/颜色/备货信息货号/成本/件数/报单名/店铺/地址/图片/网站备注/商品属性/SKC
    M列整列空白分隔
    右侧 N-S(网购): 码数/颜色/备货信息货号/图片/商品属性/SKC
  数据从第3行开始，左右同步
  C列备货信息货号=备货表SKU货号原封不动(空留空); J列网站备注=网页货号*SKU原始颜色*码数*件数
  F列固定报单名; 数据行高116; 正文宋体系列20号居中; I(左)/Q(右)列嵌SKU图140px
"""
import argparse
from copy import copy
import json
import os
import re
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from report_contract import record_key, validate_inputs, validate_report, number
from memory_store import save_json
from material_format import normalize_material_attributes
from attribute_source import attribute_signature
from local_paths import reporter_name, check_reporter

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU


def parse_args():
    p = argparse.ArgumentParser(description="TEMU报单build阶段")
    p.add_argument("--records", required=True, help="prepare输出的records.json")
    p.add_argument("--matches", required=True, help="匹配结果matches.json")
    p.add_argument("--attr-index", default=None, help="属性索引JSON(prepare的cache/product-attributes.json或by-skc-id.json)")
    p.add_argument("--output", default=None, help="输出xlsx路径；旧格式{报单名}MMDD采购单.xlsx自动改为当天第几次的名称")
    p.add_argument("--output-dir", default=None, help="自动命名报单的输出文件夹")
    return p.parse_args()


_CHINESE_NUMERALS = {
    0: "零", 1: "一", 2: "二", 3: "三", 4: "四", 5: "五",
    6: "六", 7: "七", 8: "八", 9: "九", 10: "十",
}


def chinese_number(value):
    """将正整数转为简洁中文数字，供报单文件名中的执行次数使用。"""
    if value <= 0:
        raise ValueError("执行次数必须为正整数")
    if value <= 10:
        return _CHINESE_NUMERALS[value]
    if value < 20:
        return "十" + _CHINESE_NUMERALS[value % 10]
    tens, ones = divmod(value, 10)
    result = _CHINESE_NUMERALS[tens] + "十"
    return result if ones == 0 else result + _CHINESE_NUMERALS[ones]


def _number_from_chinese(value):
    """解析当前文件名会产生的一至九十九中文次数。"""
    digits = {text: number for number, text in _CHINESE_NUMERALS.items()}
    if value in digits:
        return digits[value]
    if "十" not in value:
        return None
    left, right = value.split("十", 1)
    tens = 1 if not left else digits.get(left)
    ones = 0 if not right else digits.get(right)
    if tens is None or ones is None:
        return None
    return tens * 10 + ones


def _report_attempt_from_name(name, reporter, now):
    """从当天新旧报单文件名中提取执行次数。"""
    import re

    reporter = re.escape(str(reporter))
    readable_date = rf"{now.month}月{now.day}"
    compact_date = f"{now.month:02d}{now.day:02d}"

    current = re.fullmatch(
        rf"{reporter}{readable_date}第(?P<number>\d+|[零一二三四五六七八九十]+)次采购单\.xlsx",
        name,
    )
    if current:
        raw = current.group("number")
        return int(raw) if raw.isdigit() else _number_from_chinese(raw)

    for pattern in (
        rf"{reporter}{readable_date}采购单(?P<number>\d*)\.xlsx",
        rf"{reporter}{compact_date}采购单(?P<number>\d*)\.xlsx",
    ):
        legacy = re.fullmatch(pattern, name)
        if legacy:
            return int(legacy.group("number") or "1")
    return None


def next_daily_report_path(output_dir, reporter, now=None):
    """返回当天已有正式报单最大次数加一后的输出路径。

    同时识别中文次数、阿拉伯数字次数及旧的紧凑日期名称。只通过当前
    输出目录中的正式xlsx判断；删除文件后不会由其他缓存补回计数。
    """
    now = now or datetime.now()
    directory = Path(output_dir).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    prefix = f"{reporter}{now.month}月{now.day}第"
    attempts = [
        attempt
        for path in directory.iterdir()
        if path.is_file()
        for attempt in [_report_attempt_from_name(path.name, reporter, now)]
        if attempt is not None and attempt > 0
    ]
    attempt = max(attempts, default=0) + 1
    return directory / f"{prefix}{chinese_number(attempt)}次采购单.xlsx"


def is_legacy_auto_name(path, reporter):
    """兼容旧命令中的{报单名}MMDD采购单.xlsx，让旧调用也切换到新命名。"""
    import re
    return bool(re.fullmatch(rf"{re.escape(str(reporter))}\d{{4}}采购单(?:\d+)?\.xlsx", Path(path).name))


def resolve_reporter(matches_data, records):
    """报单名以prepare写入records的值为准；records缺失时才用matches.json或安装设置。"""
    prepared = {str(r.get("reporter") or "").strip() for r in records} - {""}
    if len(prepared) > 1:
        raise ValueError("records.json中有多个不同的报单名，请重新执行prepare。")
    prepared = next(iter(prepared), "")
    explicit = str(matches_data.get("reporter") or "").strip()
    if explicit and prepared and explicit != prepared:
        raise ValueError(
            f"matches.json的reporter（{explicit}）与prepare时的报单名（{prepared}）不一致："
            f"删除matches.json中的reporter，或用 --reporter {explicit} 重新执行prepare。"
        )
    return check_reporter(prepared or explicit or reporter_name())


class AttributeIndex(dict):
    """An explicitly supplied attribute index supersedes saved record values."""


def load_attr_index(path):
    """加载属性索引，兼容 prepare/by-skc-id/旧格式。
    返回 {键: [去重后的材料字符串,...]}（保留多个以便判断歧义，唯一才允许填）。"""
    if not path:
        return {}
    with open(path, "r", encoding="utf-8") as f:
        raw = json.load(f)
    idx = raw
    if "payload" in raw and isinstance(raw["payload"], dict) and "index" in raw["payload"]:
        directory = raw["payload"].get("details", {}).get("directoryPath")
        if directory:
            signature = attribute_signature(directory)
            if signature is None or raw.get("signature") != signature:
                raise ValueError("属性索引已过期或源目录不存在，请用当前属性库重新执行prepare后再出表。")
        idx = raw["payload"]["index"]
    elif "index" in raw:
        idx = raw["index"]
    result = AttributeIndex()
    for key, entries in idx.items():
        mats = []
        if isinstance(entries, list):
            seq = entries
        elif isinstance(entries, str):
            seq = [{"materials": entries}]
        else:
            continue
        for e in seq:
            if not isinstance(e, dict):
                continue
            m = e.get("materials", {})
            if isinstance(m, dict):
                txt = normalize_material_attributes(m)
            elif isinstance(m, str):
                txt = normalize_material_attributes(m)
            else:
                txt = ""
            txt = txt.strip()
            if txt and txt not in mats:
                mats.append(txt)
        if mats:
            result[str(key)] = mats
    return result


def _unique_material(mats):
    """材料列表唯一才返回，多个不同材料(歧义)返回空，禁止瞎选。"""
    return normalize_material_attributes(mats[0]) if len(mats) == 1 else ""


def record_attr_text(r, attr_index):
    """显式属性索引优先；没有提供索引时才用prepare保存的属性。
    有SKC只精确查；仅无SKC可按货号兜底。冲突或无结果均留空。"""
    if not isinstance(attr_index, AttributeIndex):
        if r.get("productAttributeMatch") == "ambiguous":
            return ""
        pre = (r.get("productAttributes") or "").strip()
        if pre:
            return normalize_material_attributes(pre)
    skc_id = str(r.get("skcId") or "").strip()
    if skc_id:
        return _unique_material(attr_index.get(skc_id, []))
    merged = []
    for key in (r.get("skc"), r.get("skcNumber"), r.get("rawSkuCode")):
        key = str(key or "").strip()
        if key and key in attr_index:
            for m in attr_index[key]:
                if m not in merged:
                    merged.append(m)
    return _unique_material(merged)


def build(records_path, matches_path, attr_index_path, output_path, output_dir=None):
    with open(records_path, "r", encoding="utf-8") as f:
        records = json.load(f)
    with open(matches_path, "r", encoding="utf-8") as f:
        matches_data = json.load(f)

    matches_data = validate_inputs(records, matches_data)

    reporter = resolve_reporter(matches_data, records)
    matches_data["reporter"] = reporter  # 正文与回读校验使用同一报单名
    matched = matches_data.get("matched", {})
    unmatched = matches_data.get("unmatched", [])
    pending = matches_data.get("pending", {})
    search_failed = matches_data.get("search_failed", {})

    # 默认按实际执行日命名；同日重复执行以“第一次、第二次……”顺序区分。
    # 兼容历史命令写死的{报单名}MMDD采购单.xlsx，把它当作输出目录提示而非固定文件名。
    requested_output = output_path or matches_data.get("output")
    if output_dir:
        output_path = next_daily_report_path(output_dir, reporter)
    elif not requested_output or is_legacy_auto_name(requested_output, reporter):
        directory = Path(requested_output).parent if requested_output else Path.cwd()
        output_path = next_daily_report_path(directory, reporter)
    else:
        output_path = Path(requested_output).resolve()
        if output_path.exists():
            raise FileExistsError(f"指定的输出文件已存在: {output_path}。请换一个文件名，或使用 --output-dir 自动按当天次数命名。")
    output_path = str(output_path)

    attr_index = load_attr_index(attr_index_path)
    rec_by_row = {record_key(r): r for r in records}

    # 样式（正式表以用户提供的欠货单模板为准；这里用于模板缺失时重建）
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="FFFF00")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 列定义（严格对齐模板）
    left_hdr = ["码数", "颜色", "备货信息货号", "成本", "件数", "报单名", "店铺", "地址", "图片", "网站备注", "商品属性", "SKC"]
    right_hdr = ["码数", "颜色", "备货信息货号", "图片", "商品属性", "SKC"]
    RIGHT_START = 14   # N列，右侧网购区起始列（M列=13留空分隔）
    DATA_START = 3     # 数据从第3行开始（第1行分区标题，第2行表头）
    # 模板列宽（模板缺失重建时使用）
    template_widths = {
        "A": 15.4903846153846, "B": 10.5961538461538, "C": 22.4326923076923,
        "D": 13.6153846153846, "E": 9.875, "F": 17.625, "G": 25.875,
        "H": 32.0, "I": 22.0, "J": 25.0961538461538, "K": 30.1730769230769,
        "L": 32.1442307692308, "M": 16.0, "N": 14.7019230769231,
        "O": 15.4038461538462, "P": 18.0961538461538, "Q": 20.7980769230769,
        "R": 41.6538461538462, "S": 40.4038461538462,
    }

    # 优先加载技能自带模板（保证与用户模板样式100%一致），缺失时代码重建
    skill_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    template_path = os.path.join(skill_dir, "assets", "采购单模版.xlsx")
    if os.path.exists(template_path):
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active
        ws.title = "Sheet1"
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        # 第1行分区标题（用户模板：宋体系列20号，居中，无填充）
        title_font = Font(name="Songti SC Regular", size=20, color="000000")
        title_al = Alignment(horizontal="center", vertical="center")
        ws.merge_cells("A1:L1")
        ws.merge_cells("N1:S1")
        for addr, txt in (("A1", "采购"), ("N1", "网购")):
            tc = ws[addr]
            tc.value = txt
            tc.font = title_font
            tc.alignment = title_al
        # 第2行表头（用户模板：黄色；左侧除K2商品属性外18号，K2和右侧20号）
        for i, h in enumerate(left_hdr, start=1):
            c = ws.cell(2, i, h)
            c.font = Font(name="Songti SC Regular", size=20 if i == 11 else 18)
            c.alignment = center; c.fill = hdr_fill; c.border = border
        for i, h in enumerate(right_hdr, start=RIGHT_START):
            c = ws.cell(2, i, h)
            c.font = Font(name="Songti SC Regular", size=20)
            c.alignment = center; c.fill = hdr_fill; c.border = border
        ws["M2"].fill = hdr_fill
        for col, w in template_widths.items():
            ws.column_dimensions[col].width = w
        body_font = Font(name="Songti SC Regular", size=20, color="000000")
        for row in range(DATA_START, 152):
            ws.row_dimensions[row].height = 116
            for col in list(range(1, 13)) + list(range(14, 20)):
                cell = ws.cell(row, col)
                cell.font = body_font
                cell.alignment = center
                if col not in (9, 17):
                    cell.border = border

    def force_font_size(cell, size=20):
        """只调整字号，保留模板原有字体、颜色、粗细和其他排版。"""
        font = copy(cell.font)
        font.sz = size
        cell.font = font

    def write_product_attributes(cell, value):
        """写入五行材料；两种以上鞋面主材时仅将鞋面行标红。"""
        text = str(value or "")
        force_font_size(cell)
        first_line, separator, remainder = text.partition("\n")
        _, colon, upper_materials = first_line.partition("：")
        material_count = len([part for part in re.split(r"[+、]", upper_materials) if part.strip()]) if colon else 0
        if material_count < 2:
            cell.value = text
            return

        base_font = cell.font
        red_upper_font = InlineFont(
            rFont=base_font.name or "Songti SC Regular",
            sz=base_font.sz or 20,
            color="FFFF0000",
        )
        black_font = InlineFont(
            rFont=base_font.name or "Songti SC Regular",
            sz=base_font.sz or 20,
            color="FF000000",
        )
        runs = [TextBlock(red_upper_font, first_line)]
        if separator:
            runs.append(TextBlock(black_font, "\n" + remainder))
        cell.value = CellRichText(*runs)

    # 商品属性列固定20号，模板表头和后续写入的数据都不受旧模板字号影响。
    for address in ("K2", "R2"):
        force_font_size(ws[address])

    ws.sheet_format.defaultRowHeight = 61
    ws.sheet_view.zoomScale = 60

    def ensure_template_row(excel_row):
        """Extend the row-3 template style when a batch exceeds 149 data rows."""
        if excel_row <= 151:
            return
        for col in range(1, 20):
            source = ws.cell(DATA_START, col)
            target = ws.cell(excel_row, col)
            target._style = copy(source._style)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
        ws.row_dimensions[excel_row].height = 116

    IMG_EXT = XDRPositiveSize2D(pixels_to_EMU(140), pixels_to_EMU(140))

    def add_image(path, col0, row0):
        if not path or not os.path.exists(path):
            raise ValueError(f"缺少目标图片: {path}")
        im = XLImage(path)
        marker = AnchorMarker(col=col0, row=row0)
        im.anchor = OneCellAnchor(_from=marker, ext=IMG_EXT)
        ws.add_image(im)

    # 左侧数据（从第3行开始）。已确认采购按店铺分组；待审核和搜索失败也保留
    # 在表1，方便人工继续处理，但不写入任何店铺、地址、价格或网站备注。
    exception_rows = list(pending) + list(search_failed)
    left_rows = sorted(
        list(matched) + exception_rows,
        key=lambda x: (0 if x in matched else 1, matched.get(x, {}).get("shop", ""), int(rec_by_row[x]["sourceRow"]), x),
    )
    for idx, srow in enumerate(left_rows):
        r = rec_by_row.get(srow)
        if not r:
            print(f"  [警告] 源行{srow}在records中不存在，跳过")
            continue
        excel_row = idx + DATA_START
        ensure_template_row(excel_row)
        m = matched.get(srow)
        if not m:
            sku = r.get("rawSkuCode") or ""
            skc = "（半托管订单）" if r.get("sourceType") == "semi_managed_order" else r.get("skcId") or ""
            vals = {
                1: r.get("size", ""), 2: r.get("color", ""), 3: sku, 5: r["quantity"], 6: reporter,
                11: record_attr_text(r, attr_index), 12: skc,
            }
            for col, v in vals.items():
                ws.cell(excel_row, col, v)
            write_product_attributes(ws.cell(excel_row, 11), vals[11])
            add_image(r.get("imagePath"), 8, excel_row - 1)
            ws.row_dimensions[excel_row].height = 116
            continue
        web_no = m.get("webNo", "")
        web_color = m.get("webColor", "")
        cost = m.get("cost", "")
        shop = m.get("shop", "")
        addr = m.get("address") or ""
        url = m.get("url", "")
        sku = r.get("rawSkuCode") or ""
        qty = r["quantity"]
        # 半托管订单（来自导出订单包裹表）：SKC列固定标记'（半托管订单）'
        if r.get("sourceType") == "semi_managed_order":
            skc = "（半托管订单）"
        else:
            skc = r.get("skcId") or ""
        # 规则：件数为2件时成本增加——成本列显示总成本=单价×件数（1件不变，2件翻倍）
        cost_display = float(number(cost, "单价") * number(qty, "数量"))
        vals = {
            1: r.get("size", ""), 2: r.get("color", ""), 3: sku, 4: cost_display, 5: qty,
            6: reporter, 7: shop, 8: addr,
            # 网站备注中的颜色必须原样取自本次SKU颜色；网页色名仅用于核对候选。
            10: f"{web_no}*{r.get('color', '')}*{r.get('size','')}*{qty}",
            11: record_attr_text(r, attr_index), 12: skc,
        }
        for col, v in vals.items():
            ws.cell(excel_row, col, v)
        write_product_attributes(ws.cell(excel_row, 11), vals[11])
        add_image(r.get("imagePath"), 8, excel_row - 1)
        # G列批注（规则1：白名单直接选；非白名单走综合排序，不再强制时效表对比）
        is_wl = m.get("whitelist", False)
        if m.get("memoryCopy"):
            ref = m["memoryCopy"].get("ref", "")
            url_note = url or "历史记录未填写链接"
            note = f"记忆来源: {ref}\n完整SKC命中，采购字段按历史记录照抄\n来源URL: {url_note}\n网页货号: {web_no}"
        elif is_wl:
            note = f"来源URL: {url}\n网页货号: {web_no}\n白名单店铺，鞋型核对一致直接选取"
        else:
            note = f"来源URL: {url}\n网页货号: {web_no}\n非白名单·综合排序选取"
        gc = ws.cell(excel_row, 7)
        gc.comment = openpyxl.comments.Comment(note, "report")
        gc.comment.width = 320; gc.comment.height = 130
        ws.row_dimensions[excel_row].height = 116

    # 右侧未匹配数据（N-S列，从第3行开始，与左侧同步）
    for j, srow in enumerate(unmatched):
        r = rec_by_row.get(str(srow))
        if not r:
            print(f"  [警告] 未匹配源行{srow}在records中不存在，跳过")
            continue
        excel_row = j + DATA_START
        ensure_template_row(excel_row)
        sku = r.get("rawSkuCode") or ""
        # 半托管订单：SKC列固定标记'（半托管订单）'
        if r.get("sourceType") == "semi_managed_order":
            skc = "（半托管订单）"
        else:
            skc = r.get("skcId") or ""
        # N(14)码数 O(15)颜色 P(16)备货信息货号 R(18)商品属性 S(19)SKC；Q(17)图片
        vals = {14: r.get("size", ""), 15: r.get("color", ""), 16: sku,
                18: record_attr_text(r, attr_index), 19: skc}
        for col, v in vals.items():
            ws.cell(excel_row, col, v)
        write_product_attributes(ws.cell(excel_row, 18), vals[18])
        add_image(r.get("imagePath"), 16, excel_row - 1)  # Q列=17，锚点col=16
        ws.row_dimensions[excel_row].height = 116

    # 列宽（加载模板时幂等设置一次，确保与模板一致）
    for col, w in template_widths.items():
        ws.column_dimensions[col].width = w

    # ===== Sheet2「人工审核」：仅保留网页搜款成功并已确认的记录 =====
    if "人工审核" in wb.sheetnames:
        del wb["人工审核"]
    ws2 = wb.create_sheet("人工审核")
    review_hdr = ["目标SKU图", "确定店铺商品主图", "鞋店名称"]
    review_widths = {"A": 18, "B": 22, "C": 24}
    for i, h in enumerate(review_hdr, start=1):
        c = ws2.cell(1, i, h)
        c.font = bold; c.alignment = center; c.fill = hdr_fill; c.border = border
    for col, w in review_widths.items():
        ws2.column_dimensions[col].width = w
    ws2.row_dimensions[1].height = 22
    REVIEW_EXT = XDRPositiveSize2D(pixels_to_EMU(100), pixels_to_EMU(100))

    def add_review_img(path, col0, row0):
        if not path or not os.path.exists(path):
            return
        im = XLImage(path)
        im.anchor = OneCellAnchor(_from=AnchorMarker(col=col0, row=row0), ext=REVIEW_EXT)
        ws2.add_image(im)

    # 记忆复用无需再次人工审核；待审核和搜索失败已在表1中保留但店铺信息为空。
    review_rows = sorted(
        (key for key, value in matched.items() if not value.get("memoryCopy")),
        key=lambda x: (int(rec_by_row[x]["sourceRow"]), x),
    )
    rr = 2
    for srow in review_rows:
        rec = rec_by_row.get(srow)
        if not rec:
            continue
        m = matched[srow]
        chosen = next(c for c in m.get("candidates", []) if c.get("chosen"))
        for col, value in enumerate((None, None, m["shop"]), start=1):
            cc = ws2.cell(rr, col, value); cc.border = border; cc.alignment = center
        add_review_img(rec.get("imagePath"), 0, rr - 1)
        add_review_img(chosen.get("image"), 1, rr - 1)
        for col in range(1, 4):
            ws2.cell(rr, col).border = border
        ws2.row_dimensions[rr].height = 115
        rr += 1

    fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=Path(output_path).parent)
    os.close(fd)
    try:
        wb.save(temporary)
        validation = validate_report(temporary, records, matches_data)
        os.replace(temporary, output_path)
    finally:
        if os.path.exists(temporary): os.unlink(temporary)
    print(f"已保存: {output_path}")
    print(json.dumps({"validation":validation}, ensure_ascii=False))
    # 记忆库仅允许人工确认后在原表中录入。报单程序不得新增、更新或删除记忆记录。
    memory_result = {
        "status": "manual_required",
        "report": str(Path(output_path).resolve()),
        "reason": "记忆库仅允许人工录入；本次报单未执行自动新增、更新或删除。",
    }
    receipt = Path(output_path).with_suffix(".memory-sync.json")
    save_json(receipt, memory_result)
    print(json.dumps({"memorySync":memory_result}, ensure_ascii=False))
    return output_path


def main():
    args = parse_args()
    build(args.records, args.matches, args.attr_index, args.output, args.output_dir)


if __name__ == "__main__":
    main()
