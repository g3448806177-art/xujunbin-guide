#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""TEMU 快速报单 - 标签核对辅助工具。

功能:
  1. 从标签拼图(2列N行)中裁剪单张标签，旋转到正确方向，放大
  2. 裁剪每张标签的SKC+规格区域，便于视觉识别
  3. 读取报单Excel，提取所有SKC+规格，生成核对模板
  4. 根据识别结果生成核对报告

用法:
  # 裁剪标签拼图
  python3 label_helper.py crop --image <标签拼图.jpg> --rows 4 --cols 2 --out <输出目录>

  # 放大SKC区域
  python3 label_helper.py skc --label-dir <标签目录> --out <输出目录>

  # 生成核对模板(从报单Excel提取SKC)
  python3 label_helper.py template --report <报单.xlsx> --out <核对模板.json>

  # 生成核对报告
  python3 label_helper.py report --template <核对模板.json> --identified <识别结果.json> --out <报告.md>
"""
import argparse
import json
import os
import re
import sys

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    print("需要Pillow: pip install Pillow", file=sys.stderr)
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    print("需要openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


def cmd_crop(args):
    """从拼图中裁剪单张标签"""
    im = Image.open(args.image)
    W, H = im.size
    rows, cols = args.rows, args.cols
    cw, ch = W // cols, H // rows
    os.makedirs(args.out, exist_ok=True)
    for r in range(rows):
        for c in range(cols):
            box = (c * cw + 8, r * ch + 8, (c + 1) * cw - 8, (r + 1) * ch - 8)
            crop = im.crop(box)
            # 旋转90度(标签通常是竖版，旋转后文字横向)
            crop = crop.rotate(90, expand=True)
            crop = crop.resize((int(crop.width * 1.5), int(crop.height * 1.5)), Image.LANCZOS)
            out_path = os.path.join(args.out, f"label-r{r+1}c{c+1}.jpg")
            crop.save(out_path, quality=92)
            print(f"  保存: {out_path}")
    print(f"共裁剪 {rows * cols} 张标签")


def cmd_skc(args):
    """裁剪每张标签的SKC+规格区域"""
    os.makedirs(args.out, exist_ok=True)
    labels = sorted(f for f in os.listdir(args.label_dir) if f.startswith("label-") and f.endswith(".jpg"))
    crops = []
    for fname in labels:
        im = Image.open(os.path.join(args.label_dir, fname))
        w, h = im.size
        # SKC区域在标签中下部
        crop = im.crop((int(w * 0.02), int(h * 0.40), int(w * 0.98), int(h * 0.85)))
        crop = crop.resize((int(crop.width * 1.8), int(crop.height * 1.8)), Image.LANCZOS)
        crops.append((fname, crop))
        crop.save(os.path.join(args.out, f"skc_{fname}"), quality=95)
    # 生成拼图
    if crops:
        cw, ch = crops[0][1].size
        sheet = Image.new("RGB", (cw, ch * len(crops) + 30 * len(crops)), "white")
        d = ImageDraw.Draw(sheet)
        try:
            font = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 24)
        except Exception:
            font = ImageFont.load_default()
        y = 0
        for fname, crop in crops:
            d.rectangle([0, y, cw, y + 28], fill=(40, 40, 40))
            d.text((8, y + 1), fname, fill="white", font=font)
            sheet.paste(crop, (0, y + 28))
            y += ch + 30
        sheet.thumbnail((1000, 4000), Image.LANCZOS)
        sheet.save(os.path.join(args.out, "all_skc.jpg"), quality=92)
        print(f"  拼图: {os.path.join(args.out, 'all_skc.jpg')}")
    print(f"共处理 {len(crops)} 张标签")


def cmd_template(args):
    """从报单Excel提取SKC+规格，生成核对模板"""
    wb = openpyxl.load_workbook(args.report, data_only=True)
    ws = wb.active
    items = []
    # 第1行分区标题、第2行表头，数据从第3行开始；左侧 A=码数 B=颜色 C=备货信息货号 L=SKC
    for row in range(3, ws.max_row + 1):
        size = ws.cell(row, 1).value
        color = ws.cell(row, 2).value
        sku = ws.cell(row, 3).value
        skc = ws.cell(row, 12).value
        if skc and str(skc).strip():
            items.append({
                "excelRow": row,
                "side": "left",
                "size": str(size) if size else "",
                "color": str(color) if color else "",
                "sku": str(sku) if sku else "",
                "skc": str(skc).strip(),
                "labelMatch": "",
                "labelSpec": "",
                "status": "pending",
            })
    # 右侧网购区从第3行开始（M列为空白分隔列），N=码数 O=颜色 P=备货信息货号 S=SKC
    for row in range(3, ws.max_row + 1):
        size = ws.cell(row, 14).value
        color = ws.cell(row, 15).value
        sku = ws.cell(row, 16).value
        skc = ws.cell(row, 19).value
        if skc and str(skc).strip():
            items.append({
                "excelRow": row,
                "side": "right",
                "size": str(size) if size else "",
                "color": str(color) if color else "",
                "sku": str(sku) if sku else "",
                "skc": str(skc).strip(),
                "labelMatch": "",
                "labelSpec": "",
                "status": "pending",
            })
    wb.close()
    template = {"reportPath": args.report, "itemCount": len(items), "items": items}
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(template, f, ensure_ascii=False, indent=2)
    print(f"核对模板: {args.out} ({len(items)} 条)")


def cmd_report(args):
    """根据识别结果生成核对报告"""
    with open(args.template, "r", encoding="utf-8") as f:
        template = json.load(f)
    with open(args.identified, "r", encoding="utf-8") as f:
        identified = json.load(f)  # {"labels": [{"skc": "...", "spec": "...", "position": "r1c1"}]}

    items = template["items"]
    skc_to_items = {}
    for it in items:
        skc_to_items.setdefault(it["skc"], []).append(it)

    matched = []
    spec_mismatch = []
    not_in_report = []
    for label in identified.get("labels", []):
        skc = label.get("skc", "").strip()
        spec = label.get("spec", "")
        pos = label.get("position", "")
        if skc in skc_to_items:
            found = False
            for it in skc_to_items[skc]:
                report_spec = f"{it['color']}{it['size']}"
                if spec and (spec in report_spec or report_spec in spec):
                    it["status"] = "matched"
                    it["labelMatch"] = pos
                    it["labelSpec"] = spec
                    matched.append((pos, skc, spec, report_spec))
                    found = True
                    break
            if not found:
                it = skc_to_items[skc][0]
                it["status"] = "spec_mismatch"
                it["labelMatch"] = pos
                it["labelSpec"] = spec
                spec_mismatch.append((pos, skc, spec, f"{it['color']}{it['size']}"))
        else:
            not_in_report.append((pos, skc, spec))

    # 生成报告
    lines = ["# 标签核对报告\n"]
    lines.append(f"报单: {template.get('reportPath', '')}")
    lines.append(f"报单SKC数: {len(items)}")
    lines.append(f"标签数: {len(identified.get('labels', []))}\n")
    lines.append(f"## 完全匹配: {len(matched)}\n")
    for pos, skc, spec, rs in matched:
        lines.append(f"- [{pos}] SKC={skc} 标签规格={spec} → 报单规格={rs} ✓")
    lines.append(f"\n## SKC匹配但规格不同: {len(spec_mismatch)}\n")
    for pos, skc, spec, rs in spec_mismatch:
        lines.append(f"- [{pos}] SKC={skc} 标签规格={spec} ≠ 报单规格={rs} ⚠")
    lines.append(f"\n## 报单中无此SKC: {len(not_in_report)}\n")
    for pos, skc, spec in not_in_report:
        lines.append(f"- [{pos}] SKC={skc} 规格={spec} ✗")
    # 未核对的报单项
    pending = [it for it in items if it["status"] == "pending"]
    if pending:
        lines.append(f"\n## 报单中未被标签覆盖: {len(pending)}\n")
        for it in pending:
            lines.append(f"- 行{it['excelRow']}({it['side']}) SKC={it['skc']} 规格={it['color']}{it['size']}")

    report = "\n".join(lines)
    with open(args.out, "w", encoding="utf-8") as f:
        f.write(report)
    print(f"核对报告: {args.out}")
    print(f"  完全匹配: {len(matched)}, 规格不同: {len(spec_mismatch)}, 报单中无: {len(not_in_report)}")


def main():
    parser = argparse.ArgumentParser(description="TEMU标签核对辅助工具")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_crop = sub.add_parser("crop", help="从拼图裁剪单张标签")
    p_crop.add_argument("--image", required=True)
    p_crop.add_argument("--rows", type=int, default=4)
    p_crop.add_argument("--cols", type=int, default=2)
    p_crop.add_argument("--out", required=True)

    p_skc = sub.add_parser("skc", help="放大SKC区域")
    p_skc.add_argument("--label-dir", required=True)
    p_skc.add_argument("--out", required=True)

    p_tpl = sub.add_parser("template", help="生成核对模板")
    p_tpl.add_argument("--report", required=True)
    p_tpl.add_argument("--out", required=True)

    p_rep = sub.add_parser("report", help="生成核对报告")
    p_rep.add_argument("--template", required=True)
    p_rep.add_argument("--identified", required=True)
    p_rep.add_argument("--out", required=True)

    args = parser.parse_args()
    if args.cmd == "crop":
        cmd_crop(args)
    elif args.cmd == "skc":
        cmd_skc(args)
    elif args.cmd == "template":
        cmd_template(args)
    elif args.cmd == "report":
        cmd_report(args)


if __name__ == "__main__":
    main()
