"""Shared record identity and report checks; no browser or workbook writes."""
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
import math
import re


SEARCH_FAILURE_CODES = {
    "upload_unavailable",
    "upload_not_attached",
    "upload_connection_lost",
    "search_failed",
    "rate_limited",
    "login_failed",
}


def skc_text(value):
    """Exact identifier only. No SKU/code, tail-number or fuzzy fallback."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value) if 0 <= value < 10**15 else ""
    if isinstance(value, float):
        return str(int(value)) if math.isfinite(value) and value.is_integer() and 0 <= value < 10**15 else ""
    text = str(value).strip()
    return text if re.fullmatch(r"[0-9]+", text) else ""


def record_key(record):
    return str(record.get("recordId") or record["sourceRow"])


def number(value, label, positive=False):
    if isinstance(value, bool):
        raise ValueError(f"{label}不是有效数字")
    try:
        n = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label}不是有效数字") from None
    if not n.is_finite() or n < 0 or (positive and n == 0):
        raise ValueError(f"{label}不是有效非负数字")
    return n


def normalize_matches(records, data):
    """Accept old sourceRow keys only when unambiguous across all sources."""
    ids = [record_key(r) for r in records]
    if len(ids) != len(set(ids)):
        raise ValueError("记录标识重复；多文件请重新prepare，不能按相同行号覆盖")
    by_row = {}
    for r in records:
        by_row.setdefault(str(r["sourceRow"]), []).append(record_key(r))

    def resolve(key):
        key = str(key)
        if key in ids:
            return key
        if len(by_row.get(key, [])) == 1:
            return by_row[key][0]
        raise ValueError(f"匹配键 {key} 不存在或跨文件行号有歧义，请用recordId")

    out = dict(data)
    for field in ("matched", "pending", "search_failed", "unmatchedCandidates", "searchEvidence"):
        out[field] = {}
        for key, value in (data.get(field) or {}).items():
            resolved = resolve(key)
            if resolved in out[field]:
                raise ValueError(f"{field}重复记录: {resolved}")
            out[field][resolved] = value
    out["unmatched"] = [resolve(k) for k in data.get("unmatched", [])]
    return out


def validate_inputs(records, data):
    data = normalize_matches(records, data)
    assigned = list(data["matched"]) + data["unmatched"] + list(data["pending"]) + list(data["search_failed"])
    counts = Counter(assigned)
    expected = {record_key(r) for r in records}
    if set(counts) != expected or any(n != 1 for n in counts.values()):
        raise ValueError("每条原始记录必须且只能归入匹配、网购、待审核、搜索失败中的一种")
    for r in records:
        key = record_key(r)
        qty = number(r.get("quantity"), f"{key}数量", positive=True)
        if qty != qty.to_integral_value():
            raise ValueError(f"{key}数量必须是整数")
        if not r.get("imagePath") or not Path(r["imagePath"]).is_file():
            raise ValueError(f"{key}缺少目标SKU图片")
        if key in data["matched"]:
            m = data["matched"][key]
            memory_copy = m.get("memoryCopy")
            required_fields = ("shop", "webNo", "webColor") if memory_copy else ("shop", "webNo", "webColor", "url")
            for name in required_fields:
                if not str(m.get(name) or "").strip():
                    raise ValueError(f"{key}匹配缺少{name}")
            if m.get("url") and not str(m["url"]).startswith(("https://", "http://")):
                raise ValueError(f"{key}商品链接无效")
            number(m.get("cost"), f"{key}单价")
            evidence = m.get("verification") or {}
            if memory_copy:
                if skc_text(memory_copy.get("skc")) != skc_text(r.get("skcId")) or not str(memory_copy.get("ref") or "").strip():
                    raise ValueError(f"{key}记忆复用来源不是当前完整SKC")
                if evidence.get("source") != "memory_skc_copy" or skc_text(evidence.get("skc")) != skc_text(r.get("skcId")) or not evidence.get("reason"):
                    raise ValueError(f"{key}缺少完整SKC记忆复用依据")
                if str(m.get("webColor") or "") != str(r.get("color") or ""):
                    raise ValueError(f"{key}记忆复用时webColor必须等于本次SKU颜色")
            elif any(evidence.get(k) is not True for k in ("style", "color", "size", "price")) or not evidence.get("checkedAt") or not evidence.get("reason"):
                raise ValueError(f"{key}缺少本次款式/颜色/尺码/价格核验及依据；请复核或转pending")
            chosen = [c for c in m.get("candidates", []) if c.get("chosen")]
            if not memory_copy and len(chosen) != 1:
                raise ValueError(f"{key}网页匹配必须且只能选中一个候选，供人工审核表展示确认主图")
            if len(chosen) > 1:
                raise ValueError(f"{key}选中了多个候选")
            if chosen and (chosen[0].get("shop") != m["shop"] or chosen[0].get("webNo") != m["webNo"]):
                raise ValueError(f"{key}选中候选与最终商品不一致")
            if not memory_copy and (not chosen[0].get("image") or not Path(chosen[0]["image"]).is_file()):
                raise ValueError(f"{key}缺少已确认店铺商品主图")
        elif key in data["unmatched"]:
            search = data["searchEvidence"].get(key) or {}
            memory_copy = search.get("memoryCopy")
            if memory_copy:
                if memory_copy.get("source") != "memory_skc_copy" or skc_text(memory_copy.get("skc")) != skc_text(r.get("skcId")) or not str(memory_copy.get("ref") or "").strip():
                    raise ValueError(f"{key}网购记忆复用来源不是当前完整SKC")
                if search.get("rounds"):
                    raise ValueError(f"{key}记忆复用后不得再搜图")
            else:
                rounds = search.get("rounds", [])
                wanted = ["sku", "main"] if r.get("mainImagePath") else ["sku"]
                if [x.get("kind") for x in rounds] != wanted:
                    raise ValueError(f"{key}网购记录未完成允许的搜图轮次")
                for attempt in rounds:
                    count = attempt.get("candidateCount")
                    if attempt.get("loaded") is not True or attempt.get("outcome") != "no_match" or not attempt.get("reason") or type(count) is not int or not 0 <= count <= 5:
                        raise ValueError(f"{key}搜索证据不完整；加载失败或存疑不能当作无同款")
        else:
            state = "pending" if key in data["pending"] else "search_failed"
            if not (data[state][key] or {}).get("reason"):
                raise ValueError(f"{key}的{state}缺少原因")
            if state == "search_failed":
                code = (data[state][key] or {}).get("errorCode")
                if code not in SEARCH_FAILURE_CODES:
                    allowed = "、".join(sorted(SEARCH_FAILURE_CODES))
                    raise ValueError(f"{key}的search_failed缺少有效errorCode；允许值：{allowed}")
        candidates = data["matched"].get(key, {}).get("candidates", data["unmatchedCandidates"].get(key, []))
        for c in candidates:
            if c.get("image") and not Path(c["image"]).is_file():
                raise ValueError(f"{key}候选图片文件不存在")
    return data


def validate_report(path, records, data):
    """Read back the actual workbook; verify row values, totals and image anchors."""
    import openpyxl
    data = validate_inputs(records, data)
    recs = {record_key(r): r for r in records}
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        ws, review = wb["Sheet1"], wb["人工审核"]
        matched_left = sorted(data["matched"], key=lambda k: (data["matched"][k]["shop"], int(recs[k]["sourceRow"]), k))
        exception_left = sorted(
            list(data["pending"]) + list(data["search_failed"]),
            key=lambda k: (int(recs[k]["sourceRow"]), k),
        )
        left = matched_left + exception_left
        right = data["unmatched"]
        expected_images = set()
        expected_total = Decimal(0)
        for row, key in enumerate(matched_left, 3):
            r, m = recs[key], data["matched"][key]
            total = number(m["cost"], "单价") * number(r["quantity"], "数量")
            expected_total += total
            expected = {1:r.get("size", ""), 2:r.get("color", ""), 3:r.get("rawSkuCode") or "", 4:float(total), 5:r["quantity"], 6:data.get("reporter", ""), 7:m["shop"], 8:m.get("address") or "", 10:f"{m['webNo']}*{r.get('color','')}*{r.get('size','')}*{r['quantity']}", 12:"（半托管订单）" if r.get("sourceType") == "semi_managed_order" else r.get("skcId") or ""}
            for col, value in expected.items():
                actual = ws.cell(row, col).value
                if (actual if actual is not None else "") != value:
                    raise ValueError(f"输出核对失败: {ws.cell(row,col).coordinate}")
            if not ws.cell(row,7).comment:
                raise ValueError("采购店铺缺少批注")
            expected_images.add((8,row-1))
        for row, key in enumerate(exception_left, 3 + len(matched_left)):
            r = recs[key]
            expected = {
                1:r.get("size", ""), 2:r.get("color", ""), 3:r.get("rawSkuCode") or "", 5:r["quantity"],
                6:data.get("reporter", ""),
                12:"（半托管订单）" if r.get("sourceType") == "semi_managed_order" else r.get("skcId") or "",
            }
            for col, value in expected.items():
                actual = ws.cell(row, col).value
                if (actual if actual is not None else "") != value:
                    raise ValueError(f"待审核输出核对失败: {ws.cell(row,col).coordinate}")
            if any((ws.cell(row, col).value or "") != "" for col in (4, 7, 8, 10)):
                raise ValueError("待审核行不应带有店铺或采购信息")
            if ws.cell(row, 7).comment:
                raise ValueError("待审核行不应带有店铺批注")
            expected_images.add((8,row-1))
        for row, key in enumerate(right, 3):
            r = recs[key]
            expected = {14:r.get("size", ""),15:r.get("color", ""),16:r.get("rawSkuCode") or "",19:"（半托管订单）" if r.get("sourceType") == "semi_managed_order" else r.get("skcId") or ""}
            for col,value in expected.items():
                if (ws.cell(row,col).value or "") != value:
                    raise ValueError(f"网购输出核对失败: {ws.cell(row,col).coordinate}")
            expected_images.add((16,row-1))
        actual_images = [(im.anchor._from.col,im.anchor._from.row) for im in ws._images]
        if set(actual_images) != expected_images or len(actual_images) != len(expected_images):
            raise ValueError("正式报单图片缺失、重复或错位")
        if any(ws.cell(row,13).value is not None for row in range(1,ws.max_row+1)):
            raise ValueError("M分隔列不是空白")
        for row in range(3,ws.max_row+1):
            if row >= 3+len(left) and any(ws.cell(row,col).value is not None for col in range(1,13)):
                raise ValueError("采购区存在额外数据行")
            if row >= 3+len(right) and any(ws.cell(row,col).value is not None for col in range(14,20)):
                raise ValueError("网购区存在额外数据行")
        reviewed = [k for k, m in data["matched"].items() if not m.get("memoryCopy")]
        if [review.cell(1, col).value for col in range(1, 4)] != ["目标SKU图", "确定店铺商品主图", "鞋店名称"] or review.max_column != 3:
            raise ValueError("人工审核表必须且只能有目标SKU图、确定店铺商品主图、鞋店名称三列")
        if review.max_row != len(reviewed)+1:
            raise ValueError("审核表不应包含未成功审核的记录")
        expected_review_images = {(col, row) for col in (0, 1) for row in range(1, len(reviewed)+1)}
        review_images = [(im.anchor._from.col, im.anchor._from.row) for im in review._images]
        if len(review_images) != len(expected_review_images) or set(review_images) != expected_review_images:
            raise ValueError("人工审核表的目标SKU图或确认店铺主图缺失、重复或错位")
        review_keys = sorted(reviewed, key=lambda x: (int(recs[x]["sourceRow"]), x))
        for row, key in enumerate(review_keys, start=2):
            if review.cell(row, 3).value != data["matched"][key]["shop"]:
                raise ValueError("人工审核表鞋店名称与已确认结果不一致")
        for sheet in wb:
            for row in sheet:
                if any(c.data_type == "e" for c in row):
                    raise ValueError("输出包含Excel错误单元格")
        return {"matched":len(matched_left), "unmatched":len(right), "pending":len(data["pending"]), "searchFailed":len(data["search_failed"]), "totalCost":float(expected_total)}
    finally:
        wb.close()
